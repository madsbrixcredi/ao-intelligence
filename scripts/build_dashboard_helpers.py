import csv
import os
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
PROCESSED = ROOT / "data/processed"

CANONICAL_AO_OVERRIDES = {
    "16119040": {
        "canonical_name": "BUUS JENSEN I/S STATSAUTORISEREDE REVISORER",
        "canonical_cvr": "36029374",
    }
}


def latest_master_workbook():
    explicit = os.environ.get("MASTER_WORKBOOK")
    if explicit:
        return ROOT / explicit if not Path(explicit).is_absolute() else Path(explicit)
    candidates = sorted(RAW.glob("Downloads master sheet*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    if not candidates:
        raise FileNotFoundError("No Downloads master sheet workbook found in data/raw")
    return candidates[0]


RAW_WORKBOOK = latest_master_workbook()


def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path, fieldnames, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def clean_cvr(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"\D", "", text)


def clean_text(value):
    return str(value or "").replace("\xa0", " ").strip()


def normalize_key(value):
    text = clean_text(value).casefold()
    text = re.sub(r"[^a-z0-9æøå]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def number(value):
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def canonical_ao(name, cvr_value):
    cleaned_cvr = clean_cvr(cvr_value)
    override = CANONICAL_AO_OVERRIDES.get(cleaned_cvr)
    if override:
        return override["canonical_name"], override["canonical_cvr"]
    return clean_text(name), cleaned_cvr


def declaration_lookup():
    declarations = {}
    for declaration_type, filename in [
        ("audits", "audits.csv"),
        ("reviews", "reviews.csv"),
        ("extended_reviews", "extended_reviews.csv"),
        ("assistance", "assistance.csv"),
    ]:
        for row in read_csv(PROCESSED / filename):
            name, cvr = canonical_ao(row.get("accounting_office"), row.get("accounting_office_cvr"))
            if not cvr:
                continue
            if cvr not in declarations:
                declarations[cvr] = {
                    "accounting_office_name": name,
                    "accounting_office_cvr": cvr,
                    "audits": 0,
                    "reviews": 0,
                    "extended_reviews": 0,
                    "assistance": 0,
                }
            declarations[cvr]["accounting_office_name"] = declarations[cvr]["accounting_office_name"] or name
            declarations[cvr][declaration_type] += number(row.get("declaration_count"))

    for row in declarations.values():
        row["total_declarations"] = row["audits"] + row["reviews"] + row["extended_reviews"] + row["assistance"]
    return declarations


def taxonomy_lookup():
    alias_to_canonical = {}
    cvr_to_name = {}
    path = PROCESSED / "accounting_office_taxonomy.csv"
    if not path.exists():
        return alias_to_canonical, cvr_to_name

    for row in read_csv(path):
        name, cvr = canonical_ao(row.get("canonical_accounting_office_name"), row.get("accounting_office_cvr"))
        alias = normalize_key(row.get("alias_name"))
        if alias and cvr:
            alias_to_canonical[alias] = (name, cvr)
        if cvr and name:
            cvr_to_name.setdefault(cvr, name)
    return alias_to_canonical, cvr_to_name


def workbook_rows(sheet_name):
    wb = load_workbook(RAW_WORKBOOK, data_only=True, read_only=True, keep_links=False)
    ws = wb[sheet_name]
    header = [clean_text(value) or None for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header[index]: values[index] for index in range(len(header)) if header[index]}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def build_user_name_lookup():
    lookup = defaultdict(Counter)
    for row in workbook_rows("Priority clients active compani"):
        email = clean_text(row.get("user_email")).lower()
        name = clean_text(row.get("user_name"))
        if email and name:
            lookup[email][name] += 1
    return {email: counter.most_common(1)[0][0] for email, counter in lookup.items()}


def activity_events(declarations, alias_to_canonical, cvr_to_name):
    raw_rows = workbook_rows("DA_data")

    valid_cvrs = set(declarations)
    name_to_valid_cvr = defaultdict(Counter)
    for row in raw_rows:
        raw_name = clean_text(row.get("requesting_company_name_fixed")) or clean_text(row.get("requesting_company_name"))
        analyzed_cvr = clean_cvr(row.get("analyzed_company_vat"))
        canonical_name, canonical_cvr = canonical_ao(raw_name, analyzed_cvr)
        if canonical_cvr in valid_cvrs:
            name_to_valid_cvr[normalize_key(raw_name)][canonical_cvr] += 1

    dominant_name_cvr = {
        name_key: counter.most_common(1)[0][0]
        for name_key, counter in name_to_valid_cvr.items()
        if counter
    }

    events = []
    for row in raw_rows:
        created = row.get("created")
        email = clean_text(row.get("user_email")).lower()
        if not isinstance(created, datetime) or not email:
            continue

        raw_name = clean_text(row.get("requesting_company_name_fixed")) or clean_text(row.get("requesting_company_name"))
        raw_name_key = normalize_key(raw_name)
        analyzed_cvr = clean_cvr(row.get("analyzed_company_vat"))
        canonical_name, canonical_cvr = canonical_ao(raw_name, analyzed_cvr)

        if raw_name_key in dominant_name_cvr:
            canonical_cvr = dominant_name_cvr[raw_name_key]
            canonical_name = cvr_to_name.get(canonical_cvr) or declarations.get(canonical_cvr, {}).get("accounting_office_name") or raw_name
        elif analyzed_cvr in valid_cvrs:
            canonical_cvr = analyzed_cvr
            canonical_name = cvr_to_name.get(canonical_cvr) or declarations.get(canonical_cvr, {}).get("accounting_office_name") or raw_name
        elif raw_name_key in alias_to_canonical:
            canonical_name, canonical_cvr = alias_to_canonical[raw_name_key]
        else:
            canonical_name = raw_name
            canonical_cvr = analyzed_cvr

        canonical_name, canonical_cvr = canonical_ao(canonical_name, canonical_cvr)
        client_cvr = clean_cvr(row.get("requesting_company_vat"))
        events.append(
            {
                "created": created,
                "analysis_type": clean_text(row.get("type")),
                "user_email": email,
                "accounting_office_name": canonical_name,
                "accounting_office_cvr": canonical_cvr,
                "client_name": clean_text(row.get("analyzed_company_name")),
                "client_cvr": client_cvr,
            }
        )
    return events


def first_latest(dates):
    if not dates:
        return "", ""
    return min(dates).date().isoformat(), max(dates).date().isoformat()


def build_activity_outputs():
    declarations = declaration_lookup()
    alias_to_canonical, cvr_to_name = taxonomy_lookup()
    user_names = build_user_name_lookup()
    events = activity_events(declarations, alias_to_canonical, cvr_to_name)

    by_ao = defaultdict(list)
    by_user = defaultdict(list)
    by_month = defaultdict(list)
    by_user_month = defaultdict(list)
    by_week = defaultdict(list)

    for event in events:
        cvr = event["accounting_office_cvr"]
        email = event["user_email"]
        created = event["created"]
        month_key = (cvr, created.year, created.month)
        week_start = (created.date() - timedelta(days=created.weekday())).isoformat()
        by_ao[cvr].append(event)
        by_user[(email, cvr)].append(event)
        by_month[month_key].append(event)
        by_user_month[(created.year, created.month, email, cvr)].append(event)
        by_week[(cvr, week_start)].append(event)

    ao_rows = []
    for cvr, declaration_row in declarations.items():
        event_rows = by_ao.get(cvr, [])
        first, latest = first_latest([row["created"] for row in event_rows])
        active_users = len({row["user_email"] for row in event_rows})
        distinct_clients = len({row["client_cvr"] for row in event_rows if row["client_cvr"]})
        total_analyses = len(event_rows)
        total_declarations = declaration_row["total_declarations"]
        adoption = total_analyses / total_declarations if total_declarations else 0
        ao_rows.append(
            {
                **declaration_row,
                "total_analyses": total_analyses,
                "active_users": active_users,
                "distinct_clients_analysed": distinct_clients,
                "first_activity_date": first,
                "latest_activity_date": latest,
                "analyses_per_declaration": f"{adoption:.6f}",
                "adoption_score": f"{adoption:.6f}",
            }
        )

    for cvr, event_rows in by_ao.items():
        if cvr in declarations:
            continue
        first, latest = first_latest([row["created"] for row in event_rows])
        total_analyses = len(event_rows)
        ao_rows.append(
            {
                "accounting_office_name": event_rows[0]["accounting_office_name"],
                "accounting_office_cvr": cvr,
                "audits": 0,
                "reviews": 0,
                "extended_reviews": 0,
                "assistance": 0,
                "total_declarations": 0,
                "total_analyses": total_analyses,
                "active_users": len({row["user_email"] for row in event_rows}),
                "distinct_clients_analysed": len({row["client_cvr"] for row in event_rows if row["client_cvr"]}),
                "first_activity_date": first,
                "latest_activity_date": latest,
                "analyses_per_declaration": "0.000000",
                "adoption_score": "0.000000",
            }
        )

    ao_rows.sort(key=lambda row: int(row["total_analyses"]), reverse=True)
    write_csv(
        PROCESSED / "ao_master.csv",
        [
            "accounting_office_name",
            "accounting_office_cvr",
            "audits",
            "reviews",
            "extended_reviews",
            "assistance",
            "total_declarations",
            "total_analyses",
            "active_users",
            "distinct_clients_analysed",
            "first_activity_date",
            "latest_activity_date",
            "analyses_per_declaration",
            "adoption_score",
        ],
        ao_rows,
    )

    user_rows = []
    for (email, cvr), event_rows in by_user.items():
        first, latest = first_latest([row["created"] for row in event_rows])
        user_rows.append(
            {
                "user_name": user_names.get(email, ""),
                "user_email": email,
                "accounting_office_name": event_rows[0]["accounting_office_name"],
                "accounting_office_cvr": cvr,
                "total_analyses": len(event_rows),
                "distinct_clients": len({row["client_cvr"] for row in event_rows if row["client_cvr"]}),
                "first_activity_date": first,
                "latest_activity_date": latest,
            }
        )
    user_rows.sort(key=lambda row: int(row["total_analyses"]), reverse=True)
    write_csv(
        PROCESSED / "user_master.csv",
        ["user_name", "user_email", "accounting_office_name", "accounting_office_cvr", "total_analyses", "distinct_clients", "first_activity_date", "latest_activity_date"],
        user_rows,
    )

    monthly_rows = []
    for (cvr, year, month), rows in sorted(by_month.items()):
        monthly_rows.append(
            {
                "accounting_office_name": rows[0]["accounting_office_name"],
                "accounting_office_cvr": cvr,
                "year": year,
                "month": month,
                "analyses": len(rows),
                "active_users": len({row["user_email"] for row in rows}),
                "distinct_clients": len({row["client_cvr"] for row in rows if row["client_cvr"]}),
            }
        )
    write_csv(PROCESSED / "monthly_activity.csv", ["accounting_office_name", "accounting_office_cvr", "year", "month", "analyses", "active_users", "distinct_clients"], monthly_rows)

    user_month_rows = []
    for (year, month, email, cvr), rows in sorted(by_user_month.items()):
        user_month_rows.append(
            {
                "year": year,
                "month": month,
                "user_email": email,
                "user_name": user_names.get(email, ""),
                "accounting_office_cvr": cvr,
                "accounting_office_name": rows[0]["accounting_office_name"],
                "analyses": len(rows),
            }
        )
    write_csv(PROCESSED / "monthly_user_activity.csv", ["year", "month", "user_email", "user_name", "accounting_office_cvr", "accounting_office_name", "analyses"], user_month_rows)

    weekly_rows = []
    for (cvr, week_start), rows in sorted(by_week.items()):
        weekly_rows.append(
            {
                "week_start": week_start,
                "accounting_office_name": rows[0]["accounting_office_name"],
                "accounting_office_cvr": cvr,
                "analyses": len(rows),
                "active_users": len({row["user_email"] for row in rows}),
                "distinct_clients": len({row["client_cvr"] for row in rows if row["client_cvr"]}),
            }
        )
    write_csv(PROCESSED / "weekly_activity.csv", ["week_start", "accounting_office_name", "accounting_office_cvr", "analyses", "active_users", "distinct_clients"], weekly_rows)


def build_connections_current():
    wb = load_workbook(RAW_WORKBOOK, data_only=True, read_only=True, keep_links=False)
    ws = wb["Data analysis master sheet"]

    rows = []
    for row_number in range(5, ws.max_row + 1):
        cvr = ws[f"GJ{row_number}"].value
        name = ws[f"GK{row_number}"].value
        active_connections = ws[f"GL{row_number}"].value
        purchased_companies = ws[f"GT{row_number}"].value
        if not cvr or not name:
            if row_number > 80:
                break
            continue
        canonical_name, canonical_cvr = canonical_ao(name, cvr)
        rows.append(
            {
                "accounting_office_cvr": canonical_cvr,
                "accounting_office_name": canonical_name,
                "active_erp_connections": int(active_connections or 0),
                "purchased_companies": int(purchased_companies or 0),
                "source": f"{RAW_WORKBOOK.name} Data analysis master sheet GL/GT",
                "snapshot_date": datetime.now().date().isoformat(),
            }
        )

    write_csv(
        PROCESSED / "ao_connections_current.csv",
        ["accounting_office_cvr", "accounting_office_name", "active_erp_connections", "purchased_companies", "source", "snapshot_date"],
        rows,
    )


def rebuild_opportunity_from_master():
    output = []
    for row in read_csv(PROCESSED / "ao_master.csv"):
        declarations = number(row.get("total_declarations"))
        analyses = number(row.get("total_analyses"))
        adoption = analyses / declarations if declarations else 0
        output.append(
            {
                "accounting_office_name": row["accounting_office_name"],
                "accounting_office_cvr": row["accounting_office_cvr"],
                "total_declarations": declarations,
                "total_analyses": analyses,
                "adoption_score": f"{adoption:.6f}",
                "untapped_potential": declarations - analyses,
                "potential_at_10_percent": round(declarations * 0.10),
                "potential_at_20_percent": round(declarations * 0.20),
                "potential_at_30_percent": round(declarations * 0.30),
            }
        )
    output.sort(key=lambda item: item["untapped_potential"], reverse=True)
    write_csv(PROCESSED / "ao_opportunity.csv", output[0].keys(), output)


def rebuild_implementation_risk():
    user_rows_by_cvr = defaultdict(list)
    for row in read_csv(PROCESSED / "user_master.csv"):
        user_rows_by_cvr[row["accounting_office_cvr"]].append(row)

    output = []
    for row in read_csv(PROCESSED / "ao_master.csv"):
        cvr = row["accounting_office_cvr"]
        total = number(row.get("total_analyses"))
        users = user_rows_by_cvr.get(cvr, [])
        top_three = sum(number(user.get("total_analyses")) for user in sorted(users, key=lambda user: number(user.get("total_analyses")), reverse=True)[:3])
        top_share = top_three / total if total else 0
        active_users = number(row.get("active_users"))
        risk = "LOW"
        if total == 0 or active_users <= 1 or top_share >= 0.80:
            risk = "HIGH"
        elif active_users <= 3 or top_share >= 0.60:
            risk = "MEDIUM"
        output.append(
            {
                "accounting_office_name": row["accounting_office_name"],
                "accounting_office_cvr": cvr,
                "active_users": active_users,
                "total_analyses": total,
                "top_3_user_share": f"{top_share:.6f}",
                "implementation_risk": risk,
            }
        )
    output.sort(key=lambda item: item["total_analyses"], reverse=True)
    write_csv(PROCESSED / "implementation_risk.csv", ["accounting_office_name", "accounting_office_cvr", "active_users", "total_analyses", "top_3_user_share", "implementation_risk"], output)


def main():
    print(f"Using workbook: {RAW_WORKBOOK}")
    build_connections_current()
    build_activity_outputs()
    rebuild_opportunity_from_master()
    rebuild_implementation_risk()


if __name__ == "__main__":
    main()
